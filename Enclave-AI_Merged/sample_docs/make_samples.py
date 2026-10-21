"""Generate demo documents for the Option A (inspection report) scenario.

Creates: SOP_Pump_Maintenance.docx (typed), Inspection_Report_Pump_P101.png (scan-like image)
and Inspection_Report_Pump_P101_scanned.pdf (image-only PDF -> needs OCR).
Run:  python make_samples.py
"""
import random
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = Path(__file__).parent


def make_sop():
    d = Document()
    d.add_heading("SOP-114: Centrifugal Pump Maintenance", 0)
    d.add_paragraph("Section 1. Purpose. This procedure defines inspection limits and required "
                    "maintenance actions for centrifugal pumps in the cooling water system.")
    d.add_heading("Section 2. Wear tolerance limits", 1)
    d.add_paragraph("The maximum permitted impeller wear is 3.5 mm. Impeller wear above 3.5 mm "
                    "means the impeller must be replaced before the pump returns to service.")
    d.add_paragraph("Shaft vibration must not exceed 7.1 mm/s RMS. Readings between 4.5 and 7.1 mm/s "
                    "require re-inspection within 30 days.")
    d.add_heading("Section 3. Required actions", 1)
    d.add_paragraph("If wear exceeds the limit: isolate the pump, tag out the motor, replace the "
                    "impeller and wear rings, then run a 2 hour test before release.")
    d.add_paragraph("If bearing temperature exceeds 85 degrees C the bearings must be replaced.")
    t = d.add_table(rows=3, cols=3)
    for row, vals in zip(t.rows, [("Parameter", "Limit", "Action"),
                                  ("Impeller wear", "3.5 mm", "Replace impeller"),
                                  ("Bearing temp", "85 C", "Replace bearings")]):
        for cell, v in zip(row.cells, vals):
            cell.text = v
    d.save(OUT / "SOP_Pump_Maintenance.docx")


def make_scan():
    lines = [
        "INSPECTION REPORT - PUMP P-101",
        "Date: 12 September 2026    Inspector: R. Sharma",
        "",
        "Equipment: Cooling water pump P-101",
        "Measured impeller wear: 4.2 mm",
        "Shaft vibration: 3.8 mm/s RMS",
        "Bearing temperature: 71 C",
        "Notes: Visible scoring on impeller vanes.",
        "Recommendation: refer to SOP-114.",
    ]
    img = Image.new("L", (1240, 900), 245)
    dr = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 38)
    except OSError:
        font = ImageFont.load_default()
    y = 60
    for line in lines:
        dr.text((70, y), line, fill=25, font=font)
        y += 80
    random.seed(4)
    for _ in range(2500):                                   # scanner noise
        x, yy = random.randrange(img.width), random.randrange(img.height)
        img.putpixel((x, yy), random.randrange(150, 255))
    img = img.rotate(0.6, fillcolor=245).filter(ImageFilter.GaussianBlur(0.6))
    img.save(OUT / "Inspection_Report_Pump_P101.png")
    img.convert("RGB").save(OUT / "Inspection_Report_Pump_P101_scanned.pdf", resolution=150)


if __name__ == "__main__":
    make_sop()
    make_scan()
    print("Sample documents created in", OUT)
